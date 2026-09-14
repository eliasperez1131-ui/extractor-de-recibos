from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill(start_color="FF1F4E78", end_color="FF1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="FF1F4E78")
SUMMARY_FILL = PatternFill(start_color="FFD9E1F2", end_color="FFD9E1F2", fill_type="solid")


def write_excel(output_path: str | Path, columns: List[str], rows: List[Dict[str, str]],
                locale_info: Dict | None = None, summary: Dict | None = None) -> Path:
    """Write the final Excel file with detected columns and rows."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    ws_data = wb.active
    ws_data.title = "Datos"

    ws_data.append(columns)
    for cell in ws_data[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in rows:
        ordered = []
        for col in columns:
            val = row.get(col, "")
            if val is None:
                val = ""
            ordered.append(str(val))
        ws_data.append(ordered)

    for col_idx, col_name in enumerate(columns, start=1):
        max_len = len(str(col_name))
        for row in ws_data.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx, values_only=True):
            for cell_value in row:
                if cell_value is not None:
                    max_len = max(max_len, len(str(cell_value)))
        ws_data.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

    ws_data.freeze_panes = "A2"
    ws_data.auto_filter.ref = ws_data.dimensions

    if locale_info or summary:
        ws_summary = wb.create_sheet("Resumen", 0)
        ws_summary["A1"] = "Extractor de Recibos a Excel"
        ws_summary["A1"].font = TITLE_FONT
        ws_summary.merge_cells("A1:B1")

        row = 3
        ws_summary.cell(row=row, column=1, value="Fecha de generación").fill = SUMMARY_FILL
        ws_summary.cell(row=row, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        row += 1

        if locale_info:
            ws_summary.cell(row=row, column=1, value="País/Locale detectado").fill = SUMMARY_FILL
            ws_summary.cell(row=row, column=2, value=f"{locale_info.get('country', 'N/A')} ({locale_info.get('code', 'N/A')})")
            row += 1
            ws_summary.cell(row=row, column=1, value="Moneda").fill = SUMMARY_FILL
            ws_summary.cell(row=row, column=2, value=f"{locale_info.get('currency', 'N/A')} ({locale_info.get('currency_symbol', '')})")
            row += 1
            ws_summary.cell(row=row, column=1, value="Etiqueta impuesto").fill = SUMMARY_FILL
            ws_summary.cell(row=row, column=2, value=locale_info.get("tax_label", "N/A"))
            row += 1
            ws_summary.cell(row=row, column=1, value="Etiqueta ID tributario").fill = SUMMARY_FILL
            ws_summary.cell(row=row, column=2, value=locale_info.get("tax_id_label", "N/A"))
            row += 1
            ws_summary.cell(row=row, column=1, value="Formato de fecha").fill = SUMMARY_FILL
            ws_summary.cell(row=row, column=2, value=locale_info.get("date_format", "N/A"))
            row += 1

        if summary:
            ws_summary.cell(row=row, column=1, value="Archivos procesados").fill = SUMMARY_FILL
            ws_summary.cell(row=row, column=2, value=summary.get("processed", 0))
            row += 1
            ws_summary.cell(row=row, column=1, value="PDFs (texto)").fill = SUMMARY_FILL
            ws_summary.cell(row=row, column=2, value=summary.get("pdf_text", 0))
            row += 1
            ws_summary.cell(row=row, column=1, value="PDFs (OCR)").fill = SUMMARY_FILL
            ws_summary.cell(row=row, column=2, value=summary.get("pdf_ocr", 0))
            row += 1
            ws_summary.cell(row=row, column=1, value="Imágenes (OCR)").fill = SUMMARY_FILL
            ws_summary.cell(row=row, column=2, value=summary.get("image_ocr", 0))
            row += 1
            ws_summary.cell(row=row, column=1, value="Errores").fill = SUMMARY_FILL
            ws_summary.cell(row=row, column=2, value=summary.get("errors", 0))
            row += 1
            ws_summary.cell(row=row, column=1, value="Columnas detectadas").fill = SUMMARY_FILL
            ws_summary.cell(row=row, column=2, value=len(columns))
            row += 1

        ws_summary.column_dimensions["A"].width = 30
        ws_summary.column_dimensions["B"].width = 50

    wb.save(str(output_path))
    return output_path
